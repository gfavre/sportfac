"""Printable attendance worksheets; keep the column layout independent of the view."""

from io import BytesIO

from django.utils import timezone
from dynamic_preferences.registries import global_preferences_registry
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter

from registrations.models import ChildActivityLevel

from .extra_columns import parse_extra_columns


def next_attendance_date(course):
    return next((day for day in sorted(course.all_dates) if day >= timezone.localdate()), None)


def registration_answers(registration):
    answers = {}
    for extra in registration.extra_infos.all():
        value = extra.value
        if extra.key.is_boolean or extra.key.is_image:
            value = "Oui" if extra.is_true else "Non" if extra.is_false else extra.value
        answers[extra.key.question_label] = (extra.value, value)
    return answers


def extra_column_value(column, answers):
    answer = answers.get(column["question"])
    if answer is None:
        return ""
    stored, displayed = answer
    return column.get("values", {}).get(stored, displayed)


def attendance_workbook(course, day):
    columns = parse_extra_columns(global_preferences_registry.manager()["site__ATTENDANCE_EXTRA_COLUMNS"])
    registrations = list(
        course.participants.select_related("child")
        .prefetch_related("extra_infos__key")
        .order_by("child__last_name", "child__first_name", "pk")
    )
    levels = {
        level.child_id: level
        for level in ChildActivityLevel.objects.filter(
            activity=course.activity, child_id__in=[registration.child_id for registration in registrations]
        )
    }
    headers = [
        "Absence",
        "Gp",
        "N°",
        *(column["label"] for column in columns),
        "N° cours",
        "Prénom",
        "Nom",
        "N-1",
        f"N {day.year}",
        "Vient de",
        "Va à",
    ]
    last_column = len(headers)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Absences"
    sheet.sheet_view.showGridLines = False
    sheet.append(
        [
            f"N° cours : {course.number}",
            None,
            len(registrations),
            "Responsable : " + ", ".join(person.get_full_name() for person in course.instructors.all()),
        ]
    )
    sheet.merge_cells("A1:B1")
    sheet.merge_cells(start_row=1, start_column=4, end_row=1, end_column=last_column - 2)
    date_cell = sheet.cell(1, last_column - 1, day)
    sheet.merge_cells(start_row=1, start_column=last_column - 1, end_row=1, end_column=last_column)
    date_cell.number_format = '"Date : "dd.mm.yyyy'
    sheet.append([])
    sheet.append(headers)
    for registration in registrations:
        answers = registration_answers(registration)
        level = levels.get(registration.child_id)
        sheet.append(
            [
                "",
                course.group_name,
                registration.child.bib_number,
                *(extra_column_value(column, answers) for column in columns),
                course.number,
                registration.child.first_name,
                registration.child.last_name,
                level.before_level if level else "",
                level.after_level if level else "",
                "",
                "",
            ]
        )
    format_attendance_sheet(sheet, len(columns), last_column)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def format_attendance_sheet(sheet, extra_column_count, last_column):
    sheet.auto_filter.ref = f"A3:{get_column_letter(last_column)}{max(sheet.max_row, 4)}"
    last_row = max(sheet.max_row + 3, 16)
    for row in range(4, last_row + 1):
        sheet.row_dimensions[row].height = 30
    side = Side(style="thin", color="999999")
    border = Border(left=side, right=side, top=side, bottom=side)
    for row in sheet.iter_rows(min_row=1, max_row=last_row, max_col=last_column):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Names beginning with '=' must remain text in Excel.
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.font = Font(name="Arial", size=11, bold=cell.row in (1, 3))
            if cell.row >= 3:
                cell.border = border
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="E8E8E8")
    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor="000000")
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    widths = [12, 7, 9, *([12] * extra_column_count), 12, 20, 24, 10, 10, 14, 14]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 45
    sheet.row_dimensions[3].height = 25
    sheet.print_title_rows = "1:3"
    sheet.print_area = f"A1:{get_column_letter(last_column)}{last_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
